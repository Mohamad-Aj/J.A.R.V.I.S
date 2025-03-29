from PyQt6.QtWidgets import (
    QWidget,
    QVBoxLayout,
    QLabel,
    QTextEdit,
    QTableWidget,
    QTableWidgetItem,
    QApplication,
    QPushButton,
    QHBoxLayout,
    QStackedWidget,
)
from PyQt6.QtGui import QColor
from PyQt6.QtCore import Qt
import pandas as pd
import openai
import os
from PyPDF2 import PdfReader
from dotenv import load_dotenv
from PyQt6.QtCore import pyqtSignal

load_dotenv()


class FileSummaryPopup(QWidget):
    update_summary_signal = pyqtSignal(str)
    show_error_signal = pyqtSignal(str)

    def __init__(self, file_path: str):
        super().__init__()
        # self.update_summary_signal.connect(self.summary_box.append)
        # self.show_error_signal.connect(self.summary_box.setText)
        self._stream_buffer = ""

        self.setWindowTitle("File Summary")
        self.setFixedSize(600, 400)
        self.setAttribute(Qt.WidgetAttribute.WA_TranslucentBackground)
        self.setWindowFlags(
            Qt.WindowType.FramelessWindowHint | Qt.WindowType.WindowStaysOnTopHint
        )
        self.dragging = False
        self.offset = None
        self.file_path = file_path

        self.setStyleSheet(
            """
            QWidget {
                background-color: #1e1e1e;
                color: white;
                font-size: 14px;
            }
        """
        )

        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(0, 0, 0, 0)

        title_bar = QWidget()
        title_bar.setFixedHeight(40)
        title_bar.setStyleSheet(
            """
            background-color: #2b2b2b;
            border-top-left-radius: 20px;
            border-top-right-radius: 20px;
        """
        )
        title_bar_layout = QHBoxLayout(title_bar)
        title_bar_layout.setContentsMargins(10, 0, 10, 0)
        title_bar_layout.setSpacing(12)

        self.btn_content = QPushButton("File Content")
        self.btn_summary = QPushButton("Summary")
        self.btn_close = QPushButton("✕")

        for btn in [self.btn_content, self.btn_summary]:
            btn.setCheckable(True)
            btn.setStyleSheet(
                """
                QPushButton {
                    background-color: transparent;
                    color: white;
                    padding: 6px 14px;
                    font-size: 14px;
                    border: none;
                }
                QPushButton:checked {
                    background-color: #448aff;
                    border-radius: 8px;
                }
            """
            )
            btn.setFixedHeight(26)

        self.btn_close.setFixedSize(28, 28)
        self.btn_close.setStyleSheet(
            """
            QPushButton {
                background-color: transparent;
                color: white;
                font-size: 16px;
                border: none;
            }
            QPushButton:hover {
                color: #ff5252;
            }
        """
        )
        self.btn_close.clicked.connect(self.close)

        self.btn_content.setChecked(True)
        self.btn_content.clicked.connect(lambda: self.switch_tab("content"))
        self.btn_summary.clicked.connect(lambda: self.switch_tab("summary"))

        title_bar_layout.addWidget(self.btn_content)
        title_bar_layout.addWidget(self.btn_summary)
        title_bar_layout.addStretch()
        title_bar_layout.addWidget(self.btn_close)

        main_layout.addWidget(title_bar)

        self.tabs_stack = QStackedWidget()
        self.content_tab = QWidget()
        self.summary_tab = QWidget()

        self.tabs_stack.addWidget(self.content_tab)
        self.tabs_stack.addWidget(self.summary_tab)
        main_layout.addWidget(self.tabs_stack)

        self.build_content_tab(file_path)
        self.build_summary_tab(file_path)

    def switch_tab(self, tab):
        if tab == "content":
            self.tabs_stack.setCurrentIndex(0)
            self.btn_content.setChecked(True)
            self.btn_summary.setChecked(False)
        else:
            self.tabs_stack.setCurrentIndex(1)
            self.btn_content.setChecked(False)
            self.btn_summary.setChecked(True)

    def build_content_tab(self, file_path):
        layout = QVBoxLayout()
        ext = os.path.splitext(file_path)[1].lower()

        try:
            if ext == ".csv":
                df = pd.read_csv(file_path)

            elif ext == ".xlsx":
                df = pd.read_excel(file_path)

            if ext in [".csv", ".xlsx"]:
                table = QTableWidget()
                table.setRowCount(min(50, len(df)))
                table.setColumnCount(len(df.columns))
                table.setHorizontalHeaderLabels(df.columns.tolist())
                for i in range(min(50, len(df))):
                    for j in range(len(df.columns)):
                        table.setItem(i, j, QTableWidgetItem(str(df.iat[i, j])))
                layout.addWidget(table)

            elif ext == ".pdf":
                reader = PdfReader(file_path)
                text = "\n\n".join(
                    page.extract_text() or "" for page in reader.pages[:5]
                )
                box = QTextEdit()
                box.setReadOnly(True)
                box.setText(text.strip() or "No readable text found in PDF.")
                layout.addWidget(box)

            # Inside build_content_tab:
            elif ext == ".docx":
                from docx import Document  # Add this at the top

                doc = Document(file_path)
                text = "\n\n".join(
                    [para.text for para in doc.paragraphs if para.text.strip()]
                )
                box = QTextEdit()
                box.setReadOnly(True)
                box.setText(text.strip() or "No readable text found in DOCX.")
                layout.addWidget(box)

            else:
                layout.addWidget(QLabel("Unsupported file type for content preview."))
        except Exception as e:
            layout.addWidget(QLabel(f"Error loading file: {e}"))

        self.content_tab.setLayout(layout)

    def build_summary_tab(self, file_path):
        layout = QVBoxLayout()
        self.summary_box = QTextEdit()
        self.summary_box.setReadOnly(True)
        self.summary_box.setText("Generating summary...")
        layout.addWidget(self.summary_box)
        self.summary_tab.setLayout(layout)

        # ✅ Connect signals after summary_box is defined
        self.update_summary_signal.connect(self._append_streamed_text)
        self.show_error_signal.connect(self.summary_box.setText)

        self.generate_summary(file_path)

    def _append_streamed_text(self, chunk: str):
        self._stream_buffer += chunk
        self.summary_box.setPlainText(self._stream_buffer)
        self.summary_box.verticalScrollBar().setValue(
            self.summary_box.verticalScrollBar().maximum()
        )

    def generate_summary(self, file_path):
        import threading

        def run():
            try:
                ext = os.path.splitext(file_path)[1].lower()
                sample = ""

                if ext == ".csv":
                    df = pd.read_csv(file_path)
                    sample = df.to_string(index=False)
                elif ext == ".xlsx":
                    df = pd.read_excel(file_path)
                    sample = df.to_string(index=False)
                elif ext == ".pdf":
                    reader = PdfReader(file_path)
                    sample = "\n\n".join(
                        page.extract_text() or "" for page in reader.pages[:5]
                    )
                elif ext == ".docx":
                    from docx import Document

                    doc = Document(file_path)
                    sample = "\n\n".join(
                        [para.text for para in doc.paragraphs if para.text.strip()]
                    )

                prompt = f"""
                You are a smart assistant. A user uploaded this file. Give a concise high-level summary of its contents.
                Here is the data:
                {sample}
                """

                response = openai.chat.completions.create(
                    model="gpt-4",
                    messages=[
                        {
                            "role": "system",
                            "content": "You are a data analyst assistant.",
                        },
                        {"role": "user", "content": prompt},
                    ],
                    stream=True,
                )

                for chunk in response:
                    delta = chunk.choices[0].delta
                    if delta and getattr(delta, "content", None):
                        self.update_summary_signal.emit(delta.content)

                print("[JARVIS] ✅ Finished streaming summary.")

            except Exception as e:
                self.show_error_signal.emit(f"[Error while summarizing]\n{e}")

        threading.Thread(target=run, daemon=True).start()

    def mousePressEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            self.dragging = True
            self.offset = event.globalPosition().toPoint() - self.pos()

    def mouseMoveEvent(self, event):
        if self.dragging:
            self.move(event.globalPosition().toPoint() - self.offset)

    def mouseReleaseEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            self.dragging = False
