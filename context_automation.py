import os
from PIL import Image
import pytesseract
from bidi.algorithm import get_display
from PyPDF2 import PdfReader
from docx import Document
from pptx import Presentation
import json
from openpyxl import load_workbook
import csv
import openai
import subprocess
import pyautogui
import time
import keyboard


class ContextBasedAutomation:

    def __init__(self):
        self.file_contents = {}
        pytesseract.pytesseract.tesseract_cmd = r"Tesseract-OCR\tesseract.exe"

    def organize_desktop(self):

        # First We need to gather the files from the desktop
        desktop_path = os.path.join(os.path.expanduser("~"), "Desktop")
        files = os.listdir(desktop_path)

        # Now we define the desired types
        allowed_extensions = {".py", ".c", ".h", ".txt", ".pdf", ".docx" , ".ipynb"}

        filtered_files = {}
        for file in files:
            full_path = os.path.join(desktop_path, file)
            if os.path.isfile(full_path) and any(file.endswith(ext) for ext in allowed_extensions):
                filtered_files[file] = full_path
        
        for file in filtered_files:
            print(file)

        # After saving the files and the relevant path, it's time to read them
        for file_name, file_path in filtered_files.items():
            print(f"Processing {file_name}...")
            content = self.process_file(file_path)
            if content:
                self.file_contents[file_name] = content[:2000]

        with open('output_contents.txt', 'w', encoding='utf-8') as output_file:
            for file_name, content in self.file_contents.items():
                output_file.write(f"--- Content of file: {file_name} ---\n")
                output_file.write(content + '\n')
        

        # now we want to send those files to the openai api to get categories according to the content
        all_files =self.read_txt("output_contents.txt")
        

        # Categorize using OpenAI API
        if len(all_files) == 0:
            return
        prompt = f"""You will act as a categorizer of files.
        You are provided with all the files {all_files} structured like this:
        --- Content of file: Filename.ext ---
        The file name is after "file:" so here the file name is Filename.ext. 
        After it is the content. Read all the files and give a list of categories and the file names that match each category.
        Ensure that each file belongs to one category. Additionally, avoid using generic categories like 'File Organization'. Provide meaningful categories.
        The category is a folder name that will be created later, so keep it short (maximum two words).
        """

        response = openai.chat.completions.create(
            model="gpt-3.5-turbo",
            messages=[{"role": "user", "content": prompt}],
            max_tokens=1000,
            temperature=0.2
        )
        categories = response.choices[0].message.content.strip()
        print(f"Categories: {categories}")


        current_folder = None
        for line in categories.splitlines():
            line = line.strip()
            if line and line[0].isdigit() and '.' in line:
                # Extract folder name (category name)
                folder_name = line.split('.', 1)[1].strip()
                folder_path = os.path.join(desktop_path, folder_name)

                # Create folder if it doesn't exist
                os.makedirs(folder_path, exist_ok=True)
                current_folder = folder_path
            elif line.startswith('-') and current_folder:
                # Extract file name and move it
                file_name = line.split('-', 1)[1].strip()
                source_path = os.path.join(desktop_path, file_name)

                if os.path.exists(source_path):
                    destination_path = os.path.join(current_folder, file_name)
                    os.rename(source_path, destination_path)
                    print(f"Moved {file_name} to {current_folder}")
                else:
                    print(f"File {file_name} not found on desktop.")



    # Read txt, python, c, h files
    def read_txt(self, file_path):
        with open(file_path, 'r', encoding='utf-8') as file:
            content = file.read()
            return content

    # Read text in images
    # def read_image(self, image_path):
    #     try:
    #         # Open the image
    #         img = Image.open(image_path)

    #         # Extract text from image
    #         text = pytesseract.image_to_string(img, lang='heb+eng')  # Specify 'eng' for English
    #         formatted_text = get_display(text)
    #         print(f"Extracted Text from {image_path}:\n{formatted_text}")
    #         return formatted_text
    #     except Exception as e:
    #         print(f"Error extracting text: {e}")
    #         return None

    def read_pdf(self, file_path):
        content = ""
        reader = PdfReader(file_path)
        for page in reader.pages:
            content += page.extract_text()
            content += " "
        return content

    def read_docx(self, file_path):
        content = ""
        doc = Document(file_path)
        for paragraph in doc.paragraphs:
            content += paragraph.text + "\n"
        return content

    # def read_pptx(self, file_path):
    #     content = ""
    #     presentation = Presentation(file_path)
    #     for slide in presentation.slides:
    #         for shape in slide.shapes:
    #             if shape.has_text_frame:
    #                 content += shape.text + "\n"
    #     return content

    def read_ipynb(self, file_path):
        content = ""
        with open(file_path, 'r', encoding='utf-8') as file:
            notebook = json.load(file)
            for cell in notebook.get('cells', []):
                if cell.get('cell_type') in ['code', 'markdown']:
                    content += "".join(cell.get('source', [])) + "\n"
        return content

    # def read_excel(self, file_path):
    #     workbook = load_workbook(file_path)
    #     sheet = workbook.active
    #     content = ""
    #     for row in sheet.iter_rows(values_only=True):
    #         content += str(row) + "\n"
    #     return content

    # def read_csv(self, file_path):
    #     print(f"Reading CSV file: {file_path}")
    #     content = ""
    #     try:
    #         with open(file_path, 'r', encoding='utf-8') as file:
    #             reader = csv.reader(file)
    #             for row in reader:
    #                 content += str(row) + "\n"
    #     except Exception as e:
    #         print(f"Error reading {file_path}: {e}")
    #     return content
    
    def align_desktop_icons(self):

        # Bring the desktop into focus
        keyboard.press_and_release('win + d')
        time.sleep(1)

        screen_width, screen_height = pyautogui.size()  # Get screen dimensions
        x = screen_width // 2  # Middle of the screen width
        y = screen_height // 20  # Near the top (adjusted to avoid hitting the toolbar)
        pyautogui.click(x, y)  # Simulate mouse click


        # Simulate Ctrl + Shift + F10 va
        keyboard.press('ctrl')
        keyboard.press('shift')
        keyboard.press_and_release('f10')
        keyboard.release('shift')
        keyboard.release('ctrl')
        
        time.sleep(0.5)
        keyboard.press('ctrl')
        keyboard.press('shift')
        keyboard.press_and_release('f10')
        keyboard.release('shift')
        keyboard.release('ctrl')
        time.sleep(0.5)

        # Navigate the context menu to "View > Auto Arrange Icons"
        keyboard.press_and_release('V')  
        time.sleep(0.1)
        keyboard.press_and_release('A')  
        time.sleep(0.1)
        

    def process_file(self, file_path):
        _, ext = os.path.splitext(file_path)
        ext = ext.lower()
        if ext in ['.txt', '.py', '.c', '.h']:
            return self.read_txt(file_path)
        # elif ext in ['.jpg', '.png']:
        #     return get_display(self.read_image(file_path))
        elif ext == '.pdf':
            return self.read_pdf(file_path)
        elif ext == '.docx':
            return self.read_docx(file_path)
        # elif ext == '.pptx':
        #     return self.read_pptx(file_path)
        elif ext == '.ipynb':
            return self.read_ipynb(file_path)
        # elif ext == '.xlsx':
        #     return self.read_excel(file_path)
        # elif ext == '.csv':
        #     return self.read_csv(file_path)
        else:
            print(f"Unsupported file type: {ext}")
            return None


# def main():
#     obj = ContextBasedAutomation()
#     obj.organize_desktop()
#     obj.align_desktop_icons()
#     obj.align_desktop_icons()

# if __name__ == "__main__":
#     main()
